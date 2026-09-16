"""
Profile — 简历档案（v3.1 起的全局入口）

一线求职工具（Teal / Huntr / 超级简历）的共同做法：先导入简历，其他所有功能都从这份简历派生。
本模块负责：
- 简历文本导入（粘贴 / .txt / .md / .docx）
- 规则解析成结构化档案（不依赖 LLM，秒出；有 LLM 时可再精修）
- 完整度评分 + 补全建议
- 持久化到 user_data/profile.json

其他模块读取 get_profile() 里的 raw_text / target_role / education 等字段，不再让用户重复填。
"""

import json
import os
import re
from datetime import datetime
from typing import Optional

from utils import safeCallLlm

# ═══════════════════════════════════════════════════════════
# 技能词典 — 规则层解析与 JD 匹配共用
# ═══════════════════════════════════════════════════════════

HARD_SKILLS = [
    # 编程语言
    "Python", "Java", "Go", "Golang", "C++", "C#", "C语言", "JavaScript", "TypeScript", "Kotlin", "Swift",
    "Objective-C", "Rust", "Scala", "PHP", "Ruby", "R语言", "MATLAB", "Shell", "SQL",
    # 前端 / 客户端
    "React", "Vue", "Angular", "Next.js", "Nuxt", "HTML", "CSS", "Webpack", "Vite", "Node.js", "Flutter",
    "React Native", "小程序", "Electron", "Android", "iOS", "鸿蒙", "HarmonyOS", "Unity", "Unreal",
    # 后端 / 架构
    "Spring", "Spring Boot", "Spring Cloud", "MyBatis", "Django", "Flask", "FastAPI", "Gin", "gRPC", "RESTful",
    "微服务", "分布式", "高并发", "消息队列", "Kafka", "RabbitMQ", "RocketMQ", "Redis", "MySQL", "PostgreSQL",
    "MongoDB", "Elasticsearch", "ClickHouse", "HBase", "Hive", "Spark", "Flink", "Hadoop", "Nginx",
    "Docker", "Kubernetes", "K8s", "Linux", "Git", "CI/CD", "Jenkins", "AWS", "阿里云", "腾讯云", "云原生",
    # 算法 / AI
    "机器学习", "深度学习", "PyTorch", "TensorFlow", "NLP", "自然语言处理", "计算机视觉", "CV", "推荐系统",
    "大模型", "LLM", "RAG", "Agent", "Prompt", "微调", "LangChain", "Transformer", "强化学习",
    # 数据
    "数据分析", "数据挖掘", "Excel", "Tableau", "Power BI", "PowerBI", "A/B测试", "AB测试", "埋点",
    "数据仓库", "ETL", "BI", "统计", "因果推断", "用户画像", "指标体系", "Pandas", "NumPy",
    # 产品 / 设计
    "Axure", "Figma", "Sketch", "墨刀", "原型", "PRD", "需求分析", "用户研究", "用户调研", "竞品分析",
    "产品设计", "交互设计", "UI设计", "UX", "设计系统", "Photoshop", "Illustrator", "After Effects", "C4D",
    # 运营 / 市场
    "用户运营", "内容运营", "活动运营", "社群运营", "增长", "投放", "SEO", "SEM", "私域", "直播",
    "短视频", "小红书", "抖音", "公众号", "品牌", "文案", "ROI", "GMV", "DAU", "留存", "转化率",
    # 测试 / 其他
    "自动化测试", "Selenium", "Appium", "pytest", "JMeter", "性能测试", "接口测试", "Postman",
    "英语", "CET-6", "六级", "雅思", "托福", "日语",
]

SOFT_SKILLS = [
    "沟通", "协作", "跨部门", "团队", "Owner", "owner意识", "主动", "抗压", "学习能力", "自驱",
    "数据驱动", "结果导向", "逻辑", "结构化思维", "复盘", "领导", "带团队", "项目管理", "推动",
    "用户思维", "商业化", "创新", "责任心", "细心", "表达", "汇报", "英文",
]

DEGREE_WORDS = ["博士", "硕士", "研究生", "本科", "学士", "大专", "专科", "PhD", "Master", "Bachelor", "MBA"]

SECTION_PATTERNS = {
    "education": r"^(教育背景|教育经历|教育信息|学历|教育|Education)\b",
    "projects": r"^(项目经历|项目经验|项目|个人项目|课程项目|Projects?)\b",
    "experience": r"^(实习经历|实习经验|工作经历|工作经验|实习|工作|Experience|Internships?)\b",
    "skills": r"^(专业技能|技能|技能清单|技术栈|技术能力|Skills?)\b",
    "summary": r"^(自我评价|个人评价|个人总结|个人简介|Summary|About)\b",
    "awards": r"^(荣誉奖项|获奖经历|奖项|荣誉|证书|Awards?|Honors?)\b",
    "campus": r"^(校园经历|社团经历|学生工作|校园活动)\b",
}

COMMON_ROLES = [
    "后端开发", "前端开发", "客户端开发", "算法工程师", "测试开发", "数据分析", "数据开发",
    "产品经理", "运营", "市场营销", "UI/UX设计", "游戏策划", "解决方案", "项目经理", "研发",
]


# ═══════════════════════════════════════════════════════════
# Storage
# ═══════════════════════════════════════════════════════════

def _data_dir() -> str:
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "user_data")
    os.makedirs(base, exist_ok=True)
    return os.path.abspath(base)


def _profile_path(user_id: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", user_id) or "default_user"
    return os.path.join(_data_dir(), f"profile_{safe}.json" if safe != "default_user" else "profile.json")


def empty_profile() -> dict:
    return {
        "name": "",
        "target_role": "",
        "target_cities": "",
        "job_type": "校招",          # 校招 / 社招 / 实习 / 转行
        "education": {"school": "", "major": "", "degree": "", "graduation": ""},
        "skills": [],
        "projects": [],              # [{"name", "description"}]
        "experience": [],            # [{"title", "description"}]
        "summary": "",
        "raw_text": "",
        "source": "",                # paste / txt / md / docx
        "parsed_by": "",             # rules / llm
        "created_at": "",
        "updated_at": "",
    }


def get_profile(user_id: str = "default_user") -> Optional[dict]:
    """返回档案；没有则返回 None（前端据此显示引导页）。"""
    path = _profile_path(user_id)
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not data.get("raw_text"):
            return None
        return data
    except Exception:
        return None


def save_profile(profile: dict, user_id: str = "default_user") -> dict:
    """保存档案（合并进空模板，补时间戳）。返回带完整度的档案。"""
    base = empty_profile()
    existing = get_profile(user_id) or {}
    base.update(existing)
    for k, v in (profile or {}).items():
        if k in base and v is not None:
            base[k] = v
    if isinstance(base.get("education"), dict):
        edu = empty_profile()["education"]
        edu.update(base["education"])
        base["education"] = edu
    base["skills"] = _dedupe([s.strip() for s in base.get("skills", []) if s and s.strip()])
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    base["created_at"] = base.get("created_at") or now
    base["updated_at"] = now
    with open(_profile_path(user_id), "w", encoding="utf-8") as f:
        json.dump(base, f, ensure_ascii=False, indent=2)
    return with_completeness(base)


def delete_profile(user_id: str = "default_user") -> bool:
    path = _profile_path(user_id)
    if os.path.exists(path):
        os.remove(path)
        return True
    return False


# ═══════════════════════════════════════════════════════════
# Import — 文件文本提取
# ═══════════════════════════════════════════════════════════

def extract_text_from_upload(filename: str, content: bytes) -> tuple:
    """
    从上传文件提取纯文本。支持 .txt / .md / .docx；.pdf 需要 pypdf（可选）。
    Returns: (text, source_type) ；失败时 text 为空字符串
    """
    name = (filename or "").lower()
    if name.endswith((".txt", ".md", ".markdown")):
        for enc in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return content.decode(enc), "txt" if name.endswith(".txt") else "md"
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace"), "txt"

    if name.endswith(".docx"):
        try:
            import io
            from docx import Document
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
            return "\n".join(p for p in parts if p is not None), "docx"
        except Exception:
            return "", "docx"

    if name.endswith(".pdf"):
        try:
            import io
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages), "pdf"
        except Exception:
            return "", "pdf"

    return "", "unknown"


# ═══════════════════════════════════════════════════════════
# Parse — 规则层
# ═══════════════════════════════════════════════════════════

def _dedupe(items: list) -> list:
    seen, out = set(), []
    for it in items:
        key = it.lower() if isinstance(it, str) else it
        if key not in seen:
            seen.add(key)
            out.append(it)
    return out


def _split_sections(lines: list) -> dict:
    """按常见标题把简历切成若干段。返回 {section_key: [lines]}，未识别的归入 'other'。"""
    sections = {"other": []}
    current = "other"
    for raw in lines:
        line = raw.strip().strip("【】[]#*:：-—•· ")
        matched = None
        if 0 < len(line) <= 12:
            for key, pat in SECTION_PATTERNS.items():
                if re.match(pat, line, flags=re.IGNORECASE):
                    matched = key
                    break
        if matched:
            current = matched
            sections.setdefault(current, [])
            continue
        if raw.strip():
            sections.setdefault(current, []).append(raw.rstrip())
    return sections


def _split_blocks(lines: list) -> list:
    """把一段（如项目经历）切成若干条目：标题行 + 描述。"""
    blocks, cur = [], None
    for line in lines:
        s = line.strip()
        is_bullet = bool(re.match(r"^[-•·*●▪◦]|^\d+[.、)]", s))
        looks_title = (not is_bullet) and len(s) <= 40 and not s.endswith(("。", "；", ";", ",", "，"))
        if looks_title and (cur is None or cur["description"]):
            cur = {"name": s, "description": ""}
            blocks.append(cur)
        else:
            if cur is None:
                cur = {"name": s[:30], "description": ""}
                blocks.append(cur)
            else:
                cur["description"] = (cur["description"] + "\n" + s).strip()
    return [b for b in blocks if b["name"]]


def extract_skills(text: str) -> list:
    """技能词典扫描，保持出现顺序。"""
    found = []
    lower = text.lower()
    for skill in HARD_SKILLS:
        pat = re.escape(skill.lower())
        # 英文词加边界，避免 Go 命中 Google
        if re.match(r"^[a-z0-9+#./ ]+$", skill.lower()):
            if re.search(rf"(?<![a-z0-9]){pat}(?![a-z0-9])", lower):
                found.append((lower.find(skill.lower()), skill))
        elif skill.lower() in lower:
            found.append((lower.find(skill.lower()), skill))
    found.sort(key=lambda x: x[0])
    return _dedupe([s for _, s in found])


def _extract_education(lines: list, full_text: str) -> dict:
    edu = {"school": "", "major": "", "degree": "", "graduation": ""}
    hay = "\n".join(lines) if lines else full_text
    m = re.search(r"([\u4e00-\u9fa5A-Za-z（）()·\s]{2,20}?(?:大学|学院|University|College|Institute))", hay)
    if m:
        edu["school"] = m.group(1).strip()
    for d in DEGREE_WORDS:
        if d in hay:
            edu["degree"] = {"研究生": "硕士", "学士": "本科", "专科": "大专", "Master": "硕士", "Bachelor": "本科", "PhD": "博士"}.get(d, d)
            break
    m = re.search(r"(20\d{2})[年./-]?\s*(?:毕业|届|-\s*至今|—\s*至今)?", hay)
    years = re.findall(r"20\d{2}", hay)
    if years:
        edu["graduation"] = max(years)
    # 专业：先把校名挖掉，再找带学科后缀的词，避免「华中科技大学」被当成专业
    hay_no_school = hay.replace(edu["school"], " ") if edu["school"] else hay
    for m in re.finditer(r"([\u4e00-\u9fa5]{2,14}?(?:科学与技术|工程|科学|技术|学|管理|设计|经济|金融|传播|语言|文学|法学|专业))", hay_no_school):
        cand = m.group(1).replace("专业", "").strip()
        if len(cand) < 2 or any(w in cand for w in ("大学", "学院", "学历", "学校", "教育", "本科", "硕士", "博士", "研究生")):
            continue
        edu["major"] = cand
        break
    return edu


def _guess_name(lines: list) -> str:
    for line in lines[:5]:
        s = line.strip().strip("#* ")
        if 2 <= len(s) <= 4 and re.fullmatch(r"[\u4e00-\u9fa5·]+", s):
            return s
        m = re.match(r"^(姓名|Name)[:：]\s*([\u4e00-\u9fa5A-Za-z·\s]{2,20})", s)
        if m:
            return m.group(2).strip()
    return ""


def _guess_target_role(text: str) -> str:
    m = re.search(r"(?:求职意向|意向岗位|目标岗位|应聘岗位|求职目标|Objective)[:：]?\s*([^\n，,。]{2,20})", text)
    if m:
        return m.group(1).strip()
    for role in COMMON_ROLES:
        if role in text[:400]:
            return role
    return ""


def _guess_job_type(text: str) -> str:
    if re.search(r"(工作经历|\d+年(?:以上)?(?:工作)?经验|在职|离职)", text):
        return "社招"
    if "实习" in text and re.search(r"(大[一二三四]|研[一二三]|20\d{2}届|在读)", text):
        return "校招"
    if "转行" in text:
        return "转行"
    return "校招"


def parse_resume_rules(text: str) -> dict:
    """纯规则解析。任何简历都能出结果，字段不全时靠完整度提示用户补。"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [l for l in text.split("\n")]
    sections = _split_sections(lines)

    profile = empty_profile()
    profile["raw_text"] = text.strip()
    profile["name"] = _guess_name(lines)
    profile["target_role"] = _guess_target_role(text)
    profile["job_type"] = _guess_job_type(text)
    profile["education"] = _extract_education(sections.get("education", []), text)
    profile["skills"] = extract_skills(text)
    profile["projects"] = _split_blocks(sections.get("projects", []))[:8]
    profile["experience"] = [
        {"title": b["name"], "description": b["description"]}
        for b in _split_blocks(sections.get("experience", []))[:8]
    ]
    profile["summary"] = "\n".join(sections.get("summary", []))[:600]
    profile["parsed_by"] = "rules"

    # 没识别出项目段落时，尝试从全文抓「项目」关键字附近的块，避免空档案
    if not profile["projects"]:
        cand = [l for l in lines if "项目" in l and len(l.strip()) <= 40]
        profile["projects"] = [{"name": c.strip(), "description": ""} for c in cand[:3]]
    return profile


def parse_resume_llm(text: str) -> Optional[dict]:
    """有 LLM 时的精修解析。失败返回 None，由调用方回落到规则结果。"""
    prompt = f"""把下面这份简历解析成 JSON。只输出 JSON，不要解释。

【简历原文】
{text[:8000]}

```json
{{
  "name": "姓名，没有就空字符串",
  "target_role": "求职意向岗位，没有写就根据内容推断一个最可能的",
  "job_type": "校招 / 社招 / 实习 / 转行 四选一",
  "education": {{"school": "", "major": "", "degree": "本科/硕士/博士/大专", "graduation": "毕业年份"}},
  "skills": ["技能1", "技能2"],
  "projects": [{{"name": "项目名", "description": "一段话，保留原文里的数据和结果"}}],
  "experience": [{{"title": "公司 · 岗位 · 时间", "description": "一段话，保留原文里的数据"}}],
  "summary": "一句话概括这个人（30 字内，客观，不要吹）"
}}
```

要求：不编造原文没有的数字；skills 只保留具体技能词，不要「学习能力强」这类。"""
    result = safeCallLlm(prompt, "你是一个严谨的简历结构化解析器。", output_format="json",
                         fallback={"_error": "unavailable"})
    if not isinstance(result, dict) or result.get("_error") or result.get("error") or "raw_text" in result and not result.get("skills"):
        return None
    if not any(result.get(k) for k in ("projects", "skills", "education")):
        return None
    profile = empty_profile()
    for k in ("name", "target_role", "job_type", "skills", "projects", "experience", "summary"):
        if result.get(k) is not None:
            profile[k] = result[k]
    if isinstance(result.get("education"), dict):
        profile["education"].update({k: str(v or "") for k, v in result["education"].items() if k in profile["education"]})
    profile["raw_text"] = text.strip()
    profile["parsed_by"] = "llm"
    return profile


def parse_resume(text: str, use_llm: bool = True) -> dict:
    """先规则、再（可选）LLM 精修；LLM 不可用时静默回落。"""
    rules = parse_resume_rules(text)
    if use_llm:
        llm = parse_resume_llm(text)
        if llm:
            # LLM 漏掉的技能用规则补上
            llm["skills"] = _dedupe(list(llm.get("skills", [])) + rules["skills"])
            for k in ("name", "target_role"):
                if not llm.get(k):
                    llm[k] = rules[k]
            return llm
    return rules


# ═══════════════════════════════════════════════════════════
# Completeness
# ═══════════════════════════════════════════════════════════

def completeness(profile: dict) -> dict:
    """
    完整度 0-100 + 缺什么。规则层，秒出。
    权重：原文 20 / 目标岗位 15 / 学校 15 / 项目 25 / 实习工作 10 / 技能 10 / 姓名 5
    """
    if not profile:
        return {"score": 0, "missing": ["还没有导入简历"], "level": "empty"}
    score, missing = 0, []
    raw = profile.get("raw_text", "")
    if len(raw) >= 200:
        score += 20
    elif raw:
        score += 8
        missing.append("简历原文太短（<200 字），建议粘贴完整简历")
    else:
        missing.append("没有简历原文")

    if profile.get("target_role"):
        score += 15
    else:
        missing.append("目标岗位 — 所有分析都靠它定标准")

    edu = profile.get("education") or {}
    if edu.get("school"):
        score += 10
    else:
        missing.append("学校")
    if edu.get("degree"):
        score += 5
    else:
        missing.append("学历")

    projects = profile.get("projects") or []
    if len(projects) >= 2:
        score += 25
    elif len(projects) == 1:
        score += 15
        missing.append("再补 1 个项目 — 面试官至少会追问 2 个")
    else:
        missing.append("项目经历 — 这是简历里最值钱的部分")
    if projects and not any(re.search(r"\d", p.get("description", "")) for p in projects):
        missing.append("项目描述里没有任何数字 — 没有数据的项目面试官会当课程作业")

    if profile.get("experience"):
        score += 10
    else:
        missing.append("实习 / 工作经历（没有就跳过）")

    skills = profile.get("skills") or []
    if len(skills) >= 3:
        score += 10
    elif skills:
        score += 5
        missing.append("技能只识别出 %d 个，补一下技术栈" % len(skills))
    else:
        missing.append("技能清单")

    if profile.get("name"):
        score += 5

    score = min(100, score)
    level = "high" if score >= 80 else "mid" if score >= 50 else "low"
    return {"score": score, "missing": missing, "level": level}


def with_completeness(profile: dict) -> dict:
    p = dict(profile)
    p["completeness"] = completeness(profile)
    return p


def profile_as_resume_text(profile: dict) -> str:
    """给 LLM 模块用的简历文本：优先原文，其次用结构化字段拼。"""
    if not profile:
        return ""
    if profile.get("raw_text"):
        return profile["raw_text"]
    lines = []
    if profile.get("name"):
        lines.append(profile["name"])
    edu = profile.get("education") or {}
    if any(edu.values()):
        lines.append(f"教育背景：{edu.get('school', '')} {edu.get('major', '')} {edu.get('degree', '')} {edu.get('graduation', '')}")
    if profile.get("skills"):
        lines.append("技能：" + "、".join(profile["skills"]))
    for p in profile.get("projects", []):
        lines.append(f"项目：{p.get('name', '')}\n{p.get('description', '')}")
    for e in profile.get("experience", []):
        lines.append(f"经历：{e.get('title', '')}\n{e.get('description', '')}")
    return "\n".join(lines)


def profile_summary_markdown(profile: dict) -> str:
    """工作台上的档案卡片。"""
    if not profile:
        return "> 📄 还没有简历档案。先导入简历，其他功能才会解锁。"
    c = completeness(profile)
    edu = profile.get("education") or {}
    lines = [
        f"## 📄 {profile.get('name') or '未命名'} · {profile.get('target_role') or '目标岗位未填'} · {profile.get('job_type', '')}",
        "",
        f"**完整度 {c['score']}/100**",
        "",
        f"- 教育：{edu.get('school') or '-'} · {edu.get('major') or '-'} · {edu.get('degree') or '-'}",
        f"- 技能：{'、'.join(profile.get('skills', [])[:12]) or '-'}",
        f"- 项目：{len(profile.get('projects', []))} 个 · 经历：{len(profile.get('experience', []))} 段",
    ]
    if c["missing"]:
        lines += ["", "**还缺：**"] + [f"- {m}" for m in c["missing"]]
    return "\n".join(lines)
